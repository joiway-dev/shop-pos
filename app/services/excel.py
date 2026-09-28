"""Excel export for reports. Money columns are written as numbers (Decimal
baht, never float) with a #,##0.00 format so the accountant can sum them."""

import io
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.catalog import QTY_SCALE

MONEY = "money"
QTY = "qty"
TEXT = "text"
PCT = "pct"


@dataclass
class Column:
    title: str
    kind: str = TEXT
    width: int = 14


def _cell_value(value, kind: str):
    if value is None or value == "":
        return None
    if kind == MONEY:
        return Decimal(value) / 100
    if kind == QTY:
        return Decimal(value) / QTY_SCALE
    return value


def build_xlsx(title: str, subtitle: str, columns: list[Column], rows: list[list], total: list | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws.append([])
    ws.append([c.title for c in columns])
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="FBE7D9")
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i)
        cell.font, cell.fill = Font(bold=True), fill
        ws.column_dimensions[cell.column_letter].width = c.width
    body = rows + ([total] if total else [])
    for r in body:
        ws.append([_cell_value(v, c.kind) for v, c in zip(r, columns)])
        for i, c in enumerate(columns, start=1):
            cell = ws.cell(row=ws.max_row, column=i)
            if c.kind == MONEY:
                cell.number_format = "#,##0.00"
            elif c.kind == QTY:
                cell.number_format = "#,##0.###"
    if total:
        for i in range(1, len(columns) + 1):
            ws.cell(row=ws.max_row, column=i).font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
