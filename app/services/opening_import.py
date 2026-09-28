"""Opening balances (ยอดยกมา): set each product's stock and cost from a file.

"Set", not "add": a row saying 50 makes the balance 50. The difference is
recorded as an adjust movement (ref_type "opening") carrying the cost, and
avg_cost is replayed so it equals the given cost. Preview first (rolled back),
then import the valid rows (same pattern as product import).
"""

import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Product, User
from app.models.sales import MOVE_ADJUST
from app.services import costing, stock
from app.services.audit import log_action
from app.services.money import AmountError, format_money, format_money_input, format_qty, format_qty_input, parse_money, parse_qty

COL_SKU = "รหัสสินค้า"
COL_NAME = "ชื่อสินค้า"
COL_UNIT = "หน่วยฐาน"
COL_QTY = "ยอดคงเหลือ (หน่วยฐาน)"
COL_COST = "ต้นทุนต่อหน่วยฐาน (บาท)"
HEADERS = [COL_SKU, COL_NAME, COL_UNIT, COL_QTY, COL_COST]
REQUIRED = [COL_SKU, COL_QTY]


class OpeningError(ValueError):
    pass


@dataclass
class OpeningRow:
    row_no: int
    sku: str
    name: str = ""
    before: str = ""
    after: str = ""
    cost: str = ""
    action: str = ""  # "set" | "same" | "error"
    errors: list[str] = field(default_factory=list)


@dataclass
class OpeningReport:
    rows: list[OpeningRow]
    committed: bool = False

    @property
    def changed(self) -> int:
        return sum(r.action == "set" for r in self.rows)

    @property
    def failed(self) -> int:
        return sum(r.action == "error" for r in self.rows)


def run_opening(db: Session, actor: User, rows: list[dict[str, str]], commit: bool) -> OpeningReport:
    if not actor.is_owner:
        raise OpeningError("นำเข้ายอดยกมาได้เฉพาะเจ้าของร้าน")
    results = []
    seen: set[str] = set()
    for row_no, row in enumerate(rows, start=2):
        sku = (row.get(COL_SKU) or "").strip()
        r = OpeningRow(row_no, sku)
        results.append(r)
        product = db.scalar(select(Product).where(Product.sku == sku)) if sku else None
        if product is None:
            r.action, r.errors = "error", ["ไม่พบรหัสสินค้านี้"]
            continue
        r.name = product.name
        if sku in seen:
            r.action, r.errors = "error", ["รหัสสินค้าซ้ำในไฟล์"]
            continue
        seen.add(sku)
        try:
            target = parse_qty(row.get(COL_QTY) or "", "ยอดคงเหลือ", allow_zero=True)
            cost_text = (row.get(COL_COST) or "").strip()
            cost = parse_money(cost_text, "ต้นทุน") if cost_text else product.avg_cost
        except AmountError as e:
            r.action, r.errors = "error", [str(e)]
            continue
        current = stock.balance_of(db, product.id)
        r.before, r.after, r.cost = format_qty(current), format_qty(target), format_money(cost)
        if target == current and cost == product.avg_cost:
            r.action = "same"
            continue
        stock.record_movement(db, product.id, target - current, MOVE_ADJUST, actor.id,
                              ref_type=costing.REF_OPENING, unit_cost=cost, note="ยอดยกมา")
        costing.recompute_avg_cost(db, product.id)
        r.action = "set"

    report = OpeningReport(results)
    if commit and report.changed:
        log_action(db, actor.id, "stock_opening", "stock_balances", None,
                   {"changed": report.changed, "failed": report.failed})
        db.commit()
        report.committed = True
    else:
        db.rollback()
    return report


def template_xlsx(db: Session) -> bytes:
    """Template pre-filled with every active product and its current balance/cost."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ยอดยกมา"
    ws.append(HEADERS)
    balances = stock.balances(db)
    for p in db.scalars(select(Product).where(Product.is_active.is_(True)).order_by(Product.sku)):
        ws.append([p.sku, p.name, p.base_unit, format_qty_input(balances.get(p.id, 0)), format_money_input(p.avg_cost)])
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="FBE7D9")
    for i, width in enumerate([14, 40, 10, 22, 24], start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = bold, fill
        ws.column_dimensions[cell.column_letter].width = width
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = "@"
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("วิธีกรอก")
    for line in [
        "แก้เฉพาะคอลัมน์ ยอดคงเหลือ และ ต้นทุนต่อหน่วยฐาน (คอลัมน์อื่นไว้ดูเท่านั้น)",
        "ยอดคงเหลือ = จำนวนที่นับได้จริงตอนนี้ เป็นหน่วยฐาน (เช่น ถุง) — ระบบจะ 'ตั้งยอด' ให้เท่ากับตัวเลขนี้",
        "ต้นทุนต่อหน่วยฐาน = ราคาทุนต่อ 1 หน่วยฐาน เป็นบาท (เว้นว่าง = ใช้ต้นทุนเดิม)",
        "ลบแถวของสินค้าที่ไม่ต้องการตั้งยอดออกได้",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
