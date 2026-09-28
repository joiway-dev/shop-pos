"""Product import from CSV / Excel (.xlsx).

One row per product; up to EXTRA_UNIT_COLUMNS extra units per row; aliases in
one column separated by "|". An existing SKU updates that product, otherwise a
new product is created (blank SKU = auto). Every row is run through the real
catalog service inside a savepoint, so preview and import use the same rules:
preview rolls everything back, import commits the valid rows only.
"""

import csv
import io
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Product
from app.services import catalog
from app.services.audit import log_action
from app.services.catalog import CatalogError, ProductInput, UnitInput
from app.services.matching.normalize import normalize
from app.services.money import format_money_input, format_qty_input

EXTRA_UNIT_COLUMNS = 3
MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024

COL_SKU = "รหัสสินค้า"
COL_BARCODE = "บาร์โค้ด"
COL_NAME = "ชื่อสินค้า"
COL_CATEGORY = "หมวดหมู่"
COL_BASE_UNIT = "หน่วยฐาน"
COL_BASE_PRICE = "ราคาหน่วยฐาน"
COL_DEFAULT_UNIT = "หน่วยขายหลัก"
COL_VAT_EXEMPT = "ยกเว้นVAT"
COL_MIN_STOCK = "สต็อกขั้นต่ำ"
COL_ALIASES = "ชื่อเรียก"


def unit_columns(n: int) -> tuple[str, str, str, str]:
    return (f"หน่วย{n}", f"จำนวนหน่วยฐาน{n}", f"ราคา{n}", f"บาร์โค้ด{n}")


HEADERS = [
    COL_SKU, COL_BARCODE, COL_NAME, COL_CATEGORY, COL_BASE_UNIT, COL_BASE_PRICE,
    COL_DEFAULT_UNIT, COL_VAT_EXEMPT, COL_MIN_STOCK, COL_ALIASES,
    *[c for n in range(2, EXTRA_UNIT_COLUMNS + 2) for c in unit_columns(n)],
]
REQUIRED = [COL_NAME, COL_BASE_UNIT, COL_BASE_PRICE]
TEXT_COLUMNS = {COL_SKU, COL_BARCODE, *[unit_columns(n)[3] for n in range(2, EXTRA_UNIT_COLUMNS + 2)]}

EXAMPLE_ROWS = [
    ["", "8851234567890", "ปูนซีเมนต์ตราเสือ 50 กก.", "ปูน", "ถุง", "145", "ถุง", "", "20",
     "ปูนเสือ|ปูนถุงเสือ", "พาเลท", "40", "5600", ""],
    ["", "", "เหล็กเส้นกลม 12 มม.", "เหล็ก", "เส้น", "189", "", "", "50",
     "เหล็ก 12 มิล|เหล็ก4หุน", "มัด", "10", "1850", ""],
    ["", "", "ทรายหยาบ", "ทราย-หิน", "คิว", "450", "", "ใช่", "", "ทราย", "", "", "", ""],
]

YES_VALUES = {"ใช่", "y", "yes", "true", "1", "x", "✓"}


class ImportFileError(ValueError):
    pass


@dataclass
class RowResult:
    row_no: int
    sku: str
    name: str
    action: str = ""  # "create" | "update" | "error"
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    rows: list[RowResult]
    committed: bool = False

    @property
    def created(self) -> int:
        return sum(r.action == "create" for r in self.rows)

    @property
    def updated(self) -> int:
        return sum(r.action == "update" for r in self.rows)

    @property
    def failed(self) -> int:
        return sum(r.action == "error" for r in self.rows)


# --- reading files ---------------------------------------------------------------


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "ใช่" if value else ""
    if isinstance(value, float):
        # Excel stores numbers as float; go through repr to avoid binary noise.
        d = Decimal(repr(value))
        return str(d.quantize(Decimal(1))) if d == d.to_integral_value() else format(d.normalize(), "f")
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def read_rows(filename: str, content: bytes, required: list[str] | None = None) -> list[dict[str, str]]:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("ไฟล์ใหญ่เกิน 5 MB")
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ImportFileError("เปิดไฟล์ Excel ไม่ได้ — กรุณาบันทึกเป็น .xlsx แล้วลองใหม่") from None
        raw = [[_cell_text(v) for v in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    elif name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp874"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError("อ่านไฟล์ CSV ไม่ได้ — กรุณาบันทึกเป็น CSV UTF-8")
        raw = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    else:
        raise ImportFileError("รองรับเฉพาะไฟล์ .xlsx หรือ .csv")

    raw = [r for r in raw if any(c for c in r)]
    if not raw:
        raise ImportFileError("ไฟล์ว่าง")
    header = [h.strip() for h in raw[0]]
    missing = [c for c in (required or REQUIRED) if c not in header]
    if missing:
        raise ImportFileError("ไม่พบคอลัมน์: " + ", ".join(missing) + " — กรุณาใช้ไฟล์ตัวอย่าง")
    if len(raw) - 1 > MAX_ROWS:
        raise ImportFileError(f"นำเข้าได้ครั้งละไม่เกิน {MAX_ROWS} แถว")
    return [
        {h: (row[i] if i < len(row) else "") for i, h in enumerate(header) if h}
        for row in raw[1:]
    ]


# --- import --------------------------------------------------------------------


def _row_input(db: Session, actor_id: int, row: dict[str, str], product: Product | None) -> ProductInput:
    def value(col: str, current: str = "") -> str:
        # When updating, a column missing from the file keeps the current value.
        return row[col] if col in row else current

    category_name = row.get(COL_CATEGORY, "").strip()
    category_id = catalog.get_or_create_category(db, actor_id, category_name).id if category_name else None
    if product is not None and COL_CATEGORY not in row:
        category_id = product.category_id
    units = [
        UnitInput(
            unit_name=row.get(cols[0], ""), factor=row.get(cols[1], ""),
            price=row.get(cols[2], ""), barcode=row.get(cols[3], ""),
        )
        for cols in (unit_columns(n) for n in range(2, EXTRA_UNIT_COLUMNS + 2))
    ]
    if product is not None:
        # Units not mentioned in the file are kept as they are.
        named = {u.unit_name.strip() for u in units}
        for u in product.units:
            if u.unit_name != product.base_unit and u.unit_name not in named:
                units.append(UnitInput(u.unit_name, format_qty_input(u.factor_to_base),
                                       format_money_input(u.price), u.barcode or ""))
    current_default = ""
    current_vat = ""
    current_min = "0"
    current_barcode = ""
    if product is not None:
        current_default = product.default_unit.unit_name if product.default_unit else ""
        current_vat = "ใช่" if product.vat_exempt else ""
        current_min = format_qty_input(product.min_stock_qty)
        current_barcode = product.barcode or ""
    return ProductInput(
        sku=row.get(COL_SKU, ""),
        barcode=value(COL_BARCODE, current_barcode),
        name=row.get(COL_NAME, ""),
        category_id=category_id,
        base_unit=row.get(COL_BASE_UNIT, ""),
        base_price=row.get(COL_BASE_PRICE, ""),
        default_unit=value(COL_DEFAULT_UNIT, current_default),
        vat_exempt=value(COL_VAT_EXEMPT, current_vat).strip().lower() in YES_VALUES,
        min_stock=value(COL_MIN_STOCK, current_min) or "0",
        units=units,
    )


def _flatten(errors: dict) -> list[str]:
    out = []
    for value in errors.values():
        out.extend(value if isinstance(value, list) else [value])
    return out


def run_import(db: Session, actor_id: int, rows: list[dict[str, str]], commit: bool) -> ImportReport:
    """Validate (commit=False) or import (commit=True) the rows."""
    results: list[RowResult] = []
    for row_no, row in enumerate(rows, start=2):  # row 1 = header
        sku = row.get(COL_SKU, "").strip()
        result = RowResult(row_no, sku, row.get(COL_NAME, "").strip())
        savepoint = db.begin_nested()
        try:
            product = catalog.get_product(db, pid) if (
                sku and (pid := db.scalar(select(Product.id).where(Product.sku == sku)))
            ) else None
            data = _row_input(db, actor_id, row, product)
            if product is None:
                same_name = db.scalar(
                    select(Product.sku).where(Product.name_normalized == normalize(data.name))
                )
                product = catalog.create_product(db, actor_id, data)
                result.action = "create"
                if same_name:
                    result.warnings.append(
                        f"มีสินค้าชื่อนี้อยู่แล้ว (รหัส {same_name}) — ถ้าต้องการแก้ไขสินค้าเดิมให้ใส่รหัสสินค้า"
                    )
            else:
                catalog.update_product(db, actor_id, product, data)
                result.action = "update"
            result.sku = product.sku
            existing_aliases = {a.alias_normalized for a in product.aliases}
            for text in row.get(COL_ALIASES, "").split("|"):
                if not text.strip():
                    continue
                try:
                    alias, warning = catalog.add_alias(db, actor_id, product, text)
                    existing_aliases.add(alias.alias_normalized)
                    if warning:
                        result.warnings.append(warning)
                except CatalogError as e:
                    # Already present / same as name: not worth failing the row.
                    result.warnings.append(f"ชื่อเรียก \"{text.strip()}\": {e}")
            savepoint.commit()
        except CatalogError as e:
            savepoint.rollback()
            result.action = "error"
            result.errors = _flatten(e.errors)
        results.append(result)

    report = ImportReport(results)
    if commit and (report.created or report.updated):
        log_action(db, actor_id, "product_import", "product", None,
                   {"created": report.created, "updated": report.updated, "failed": report.failed})
        db.commit()
        report.committed = True
    else:
        db.rollback()
    return report


# --- template ------------------------------------------------------------------


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "สินค้า"
    ws.append(HEADERS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDE8F7")
    for i, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = bold, fill
        ws.column_dimensions[cell.column_letter].width = max(12, len(header) + 4)
        if header in TEXT_COLUMNS:
            # Keep leading zeros of barcodes / SKUs.
            for r in range(2, 1001):
                ws.cell(row=r, column=i).number_format = "@"
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["J"].width = 28
    ws.freeze_panes = "A2"

    help_ws = wb.create_sheet("วิธีกรอก")
    for line in [
        ["คอลัมน์ที่ต้องกรอก: ชื่อสินค้า, หน่วยฐาน, ราคาหน่วยฐาน"],
        ["รหัสสินค้า: เว้นว่าง = ระบบสร้างให้ (P00001...) / ถ้าตรงกับสินค้าที่มีอยู่ = แก้ไขสินค้านั้น"],
        ["หน่วยฐาน: หน่วยที่เล็กที่สุดที่นับสต็อก เช่น ถุง, เส้น, คิว"],
        ["หน่วย2-4: หน่วยขายเพิ่ม เช่น พาเลท / จำนวนหน่วยฐาน = 1 พาเลทมีกี่ถุง (40) / ราคาต่อหน่วยนั้น"],
        ["หน่วยขายหลัก: หน่วยที่ขึ้นอัตโนมัติตอนขาย (เว้นว่าง = หน่วยฐาน)"],
        ["ชื่อเรียก: คำที่ลูกค้า/พนักงานเรียก คั่นหลายคำด้วย | เช่น ปูนเสือ|ปูนถุงเสือ"],
        ["ยกเว้นVAT: ใส่ ใช่ ถ้าสินค้าไม่ต้องเสีย VAT"],
        ["ราคา: ใส่ตัวเลขบาท เช่น 145 หรือ 145.50 (ไม่ต้องใส่ ฿)"],
        ["บาร์โค้ด: ช่องนี้ตั้งเป็นข้อความไว้แล้ว เลข 0 ข้างหน้าจะไม่หาย"],
    ]:
        help_ws.append(line)
    help_ws.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    writer.writerows(EXAMPLE_ROWS)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel shows Thai correctly
