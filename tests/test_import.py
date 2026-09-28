import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import AuditLog, Category, Product
from app.services import product_import as imp
from app.services.matching import matcher
from tests.factories import make_product


def to_csv(rows: list[list[str]], header=imp.HEADERS, encoding="utf-8-sig") -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode(encoding)


def row(**kw) -> list[str]:
    values = {h: "" for h in imp.HEADERS}
    values.update(kw)
    return [values[h] for h in imp.HEADERS]


def product_count(db) -> int:
    return db.scalar(select(func.count()).select_from(Product))


def test_template_roundtrip_xlsx(db, owner):
    content = imp.template_xlsx()
    wb = load_workbook(io.BytesIO(content))
    assert [c.value for c in wb.worksheets[0][1]] == imp.HEADERS
    rows = imp.read_rows("template.xlsx", content)
    report = imp.run_import(db, owner.id, rows, commit=True)
    assert (report.created, report.failed) == (3, 0), [r.errors for r in report.rows]
    tiger = db.scalars(select(Product).where(Product.barcode == "8851234567890")).one()
    assert {u.unit_name: (u.factor_to_base, u.price) for u in tiger.units} == {
        "ถุง": (1000, 14500), "พาเลท": (40000, 560000),
    }
    assert {a.alias_text for a in tiger.aliases} == {"ปูนเสือ", "ปูนถุงเสือ"}
    assert tiger.min_stock_qty == 20_000
    sand = db.scalars(select(Product).where(Product.name == "ทรายหยาบ")).one()
    assert sand.vat_exempt is True
    assert db.scalar(select(func.count()).select_from(Category)) == 3
    assert matcher.match(db, "ปูนถุงเสือ").best.product_id == tiger.id


def test_template_csv_readable(db, owner):
    rows = imp.read_rows("x.csv", imp.template_csv())
    assert len(rows) == 3
    assert rows[0][imp.COL_NAME] == "ปูนซีเมนต์ตราเสือ 50 กก."


def test_csv_cp874_encoding(db, owner):
    content = to_csv([row(**{imp.COL_NAME: "ทรายหยาบ", imp.COL_BASE_UNIT: "คิว", imp.COL_BASE_PRICE: "450"})],
                     encoding="cp874")
    assert imp.read_rows("x.csv", content)[0][imp.COL_NAME] == "ทรายหยาบ"


def test_preview_does_not_write(db, owner):
    rows = imp.read_rows("x.csv", to_csv([
        row(**{imp.COL_NAME: "ทรายหยาบ", imp.COL_CATEGORY: "ทราย", imp.COL_BASE_UNIT: "คิว", imp.COL_BASE_PRICE: "450"}),
    ]))
    report = imp.run_import(db, owner.id, rows, commit=False)
    assert report.created == 1 and not report.committed
    assert product_count(db) == 0
    assert db.scalar(select(func.count()).select_from(Category)) == 0


def test_invalid_rows_reported_and_skipped(db, owner):
    rows = imp.read_rows("x.csv", to_csv([
        row(**{imp.COL_NAME: "ทรายหยาบ", imp.COL_BASE_UNIT: "คิว", imp.COL_BASE_PRICE: "450"}),
        row(**{imp.COL_NAME: "", imp.COL_BASE_UNIT: "คิว", imp.COL_BASE_PRICE: "450"}),
        row(**{imp.COL_NAME: "หินคลุก", imp.COL_BASE_UNIT: "คิว", imp.COL_BASE_PRICE: "สี่ร้อย"}),
        row(**{imp.COL_NAME: "ปูน", imp.COL_BASE_UNIT: "ถุง", imp.COL_BASE_PRICE: "100",
               "หน่วย2": "พาเลท", "จำนวนหน่วยฐาน2": "0", "ราคา2": "10"}),
    ]))
    report = imp.run_import(db, owner.id, rows, commit=True)
    assert [r.action for r in report.rows] == ["create", "error", "error", "error"]
    assert [r.row_no for r in report.rows] == [2, 3, 4, 5]
    assert report.rows[2].errors and "ราคา" in report.rows[2].errors[0]
    assert product_count(db) == 1


def test_duplicate_barcode_within_file(db, owner):
    rows = imp.read_rows("x.csv", to_csv([
        row(**{imp.COL_NAME: "ก", imp.COL_BASE_UNIT: "ชิ้น", imp.COL_BASE_PRICE: "1", imp.COL_BARCODE: "0012"}),
        row(**{imp.COL_NAME: "ข", imp.COL_BASE_UNIT: "ชิ้น", imp.COL_BASE_PRICE: "1", imp.COL_BARCODE: "0012"}),
    ]))
    report = imp.run_import(db, owner.id, rows, commit=True)
    assert [r.action for r in report.rows] == ["create", "error"]
    assert db.scalars(select(Product.barcode)).one() == "0012"  # leading zeros kept


def test_update_by_sku_keeps_missing_columns_and_units(db, owner):
    p = make_product(db, owner.id, "ปูนเสือ", "ถุง", "145", sku="S1", barcode="885",
                     units=[("พาเลท", "40", "5600")])
    db.commit()
    header = [imp.COL_SKU, imp.COL_NAME, imp.COL_BASE_UNIT, imp.COL_BASE_PRICE]
    rows = imp.read_rows("x.csv", to_csv([["S1", "ปูนเสือ 50 กก", "ถุง", "150"]], header=header))
    report = imp.run_import(db, owner.id, rows, commit=True)
    assert report.updated == 1
    db.refresh(p)
    assert (p.name, p.barcode) == ("ปูนเสือ 50 กก", "885")
    assert {u.unit_name: u.price for u in p.units} == {"ถุง": 15000, "พาเลท": 560000}
    assert db.scalar(select(func.count()).where(AuditLog.action == "price_change")) == 1


def test_same_name_without_sku_warns(db, owner):
    make_product(db, owner.id, "ทรายหยาบ", "คิว", "450")
    db.commit()
    rows = imp.read_rows("x.csv", to_csv([
        row(**{imp.COL_NAME: "ทราย หยาบ", imp.COL_BASE_UNIT: "คิว", imp.COL_BASE_PRICE: "450"}),
    ]))
    report = imp.run_import(db, owner.id, rows, commit=False)
    assert report.rows[0].action == "create"
    assert any("มีสินค้าชื่อนี้อยู่แล้ว" in w for w in report.rows[0].warnings)


def test_xlsx_numeric_cells(db, owner):
    wb = Workbook()
    ws = wb.active
    ws.append(imp.HEADERS)
    values = {h: None for h in imp.HEADERS}
    values.update({imp.COL_NAME: "สายไฟ", imp.COL_BASE_UNIT: "เมตร", imp.COL_BASE_PRICE: 12.5,
                   imp.COL_BARCODE: 8851234567890, "หน่วย2": "ม้วน", "จำนวนหน่วยฐาน2": 100, "ราคา2": 1100})
    ws.append([values[h] for h in imp.HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    report = imp.run_import(db, owner.id, imp.read_rows("x.xlsx", buf.getvalue()), commit=True)
    assert report.created == 1, report.rows[0].errors
    p = db.scalars(select(Product)).one()
    assert p.barcode == "8851234567890"
    assert {u.unit_name: u.price for u in p.units} == {"เมตร": 1250, "ม้วน": 110000}


@pytest.mark.parametrize(
    "filename,content,message",
    [
        ("x.txt", b"abc", "xlsx"),
        ("x.csv", b"", "ว่าง"),
        ("x.csv", "ชื่อ,ราคา\nก,1".encode(), "ไม่พบคอลัมน์"),
        ("x.xlsx", b"not a zip", "Excel"),
    ],
)
def test_bad_files(filename, content, message):
    with pytest.raises(imp.ImportFileError) as e:
        imp.read_rows(filename, content)
    assert message in str(e.value)
