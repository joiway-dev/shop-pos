from sqlalchemy import select

from app.models import Product, Purchase
from app.services import auth, opening_import, stock
from tests.conftest import CASHIER_PIN, OWNER_PIN, login
from tests.factories import seed_catalog


def unit_id(product, name):
    return next(u.id for u in product.units if u.unit_name == name)


def test_owner_receives_goods_via_pages(client, db, owner):
    p = seed_catalog(db, owner.id)
    login(client, owner, OWNER_PIN)
    assert client.post("/suppliers", data={"name": "ปูนดี"}, follow_redirects=False).status_code == 303
    assert client.get("/purchases/new").status_code == 200
    units = client.get(f"/purchases/units/{p['tiger50'].id}").json()
    assert units["avg_cost_text"] == "0.00" and {u["unit_name"] for u in units["units"]} == {"ถุง", "พาเลท"}
    r = client.post("/purchases", json={"lines": [{"product_id": p["tiger50"].id,
                                                   "unit_id": unit_id(p["tiger50"], "พาเลท"),
                                                   "qty": "1", "unit_cost": "5200"}]}).json()
    assert r["ok"] and r["doc_no"].startswith("GR")
    page = client.get(f"/purchases/{r['id']}")
    assert "5,200.00" in page.text and "130.00" in page.text
    assert r["doc_no"] in client.get("/purchases").text
    db.rollback()
    db.expire_all()  # see what the app committed
    assert db.get(Product, p["tiger50"].id).avg_cost == 13000


def test_staff_flow_way_a(client, db, owner, cashier):
    p = seed_catalog(db, owner.id)
    login(client, cashier, CASHIER_PIN)
    assert client.get("/purchases/new").status_code == 403
    auth.set_permissions(db, owner, cashier, can_receive_stock=True, can_see_cost=False)
    db.commit()
    assert client.get("/purchases/new").status_code == 200
    assert client.get(f"/purchases/units/{p['sand'].id}").json()["avg_cost_text"] is None
    body = {"lines": [{"product_id": p["sand"].id, "unit_id": unit_id(p["sand"], "คิว"), "qty": "3"}]}
    r = client.post("/purchases", json=body).json()
    assert r["ok"]
    detail = client.get(f"/purchases/{r['id']}").text
    assert "รอเจ้าของร้านใส่ต้นทุน" in detail and "ราคาทุน" not in detail
    # Staff can't post costs.
    line_id = db.scalar(select(Purchase).where(Purchase.id == r["id"])).lines[0].id
    assert client.post(f"/purchases/{r['id']}/costs", data={f"cost_{line_id}": "1"}).status_code == 400

    login(client, owner, OWNER_PIN)
    assert "รอใส่ต้นทุน" in client.get("/purchases").text
    resp = client.post(f"/purchases/{r['id']}/costs", data={f"cost_{line_id}": "300", "vat_type": "none"},
                       follow_redirects=False)
    assert resp.status_code == 303
    db.rollback()
    db.expire_all()  # see what the app committed
    assert db.get(Purchase, r["id"]).status == "completed"
    assert db.get(Product, p["sand"].id).avg_cost == 30000


def test_void_purchase_page(client, db, owner):
    p = seed_catalog(db, owner.id)
    login(client, owner, OWNER_PIN)
    body = {"lines": [{"product_id": p["sand"].id, "unit_id": unit_id(p["sand"], "คิว"), "qty": "3", "unit_cost": "300"}]}
    r = client.post("/purchases", json=body).json()
    assert client.post(f"/purchases/{r['id']}/void", data={"reason": "รับผิด", "owner_pin": "0000"}).status_code == 400
    assert client.post(f"/purchases/{r['id']}/void", data={"reason": "รับผิด", "owner_pin": OWNER_PIN},
                       follow_redirects=False).status_code == 303
    db.rollback()
    db.expire_all()  # see what the app committed
    assert stock.balance_of(db, p["sand"].id) == 0


def test_stock_pages(client, db, owner, cashier):
    p = seed_catalog(db, owner.id)
    p["sand"].min_stock_qty = 5000
    db.commit()
    login(client, cashier, CASHIER_PIN)
    page = client.get("/stock")
    assert page.status_code == 200 and "ทรายหยาบ" in page.text and "มูลค่า" not in page.text
    assert client.get(f"/stock/{p['sand'].id}").status_code == 200
    assert client.post(f"/stock/{p['sand'].id}/count", data={"counted": "1", "reason": "นับ"}).status_code == 403
    assert client.get("/stock/opening").status_code == 403

    login(client, owner, OWNER_PIN)
    assert "มูลค่า" in client.get("/stock").text
    r = client.post(f"/stock/{p['sand'].id}/count", data={"counted": "8", "reason": "นับสต็อก"}, follow_redirects=False)
    assert r.status_code == 303
    card = client.get(f"/stock/{p['sand'].id}").text
    assert "ปรับตามนับจริง" in card and "นับสต็อก" in card
    r = client.post(f"/stock/{p['sand'].id}/writeoff", data={"qty": "1", "kind": "damaged", "reason": "เปียกน้ำ"},
                    follow_redirects=False)
    assert r.status_code == 303
    db.rollback()
    db.expire_all()  # see what the app committed
    assert stock.balance_of(db, p["sand"].id) == 7000


def test_opening_pages(client, db, owner):
    p = seed_catalog(db, owner.id)
    login(client, owner, OWNER_PIN)
    content = client.get("/stock/opening/template.xlsx").content
    r = client.post("/stock/opening", files={"file": ("o.xlsx", content, "application/octet-stream")})
    assert r.status_code == 200 and "ตรวจสอบ" in r.text
    assert client.post("/stock/opening", files={"file": ("o.csv", b"x", "text/csv")}).status_code == 400
    # Edit one row and import it.
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == p["sand"].sku:
            row[3].value, row[4].value = "12", "320"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/stock/opening", files={"file": ("o.xlsx", buf.getvalue(), "application/octet-stream")})
    token = r.text.split('name="token" value="')[1].split('"')[0]
    done = client.post("/stock/opening/confirm", data={"token": token})
    assert "ตั้งยอดเรียบร้อย 1 รายการ" in done.text
    db.rollback()
    db.expire_all()  # see what the app committed
    assert stock.balance_of(db, p["sand"].id) == 12_000
    assert db.get(Product, p["sand"].id).avg_cost == 32000


def test_users_permission_toggle(client, db, owner, cashier):
    login(client, owner, OWNER_PIN)
    r = client.post(f"/users/{cashier.id}/permissions", data={"can_receive_stock": "true"}, follow_redirects=False)
    assert r.status_code == 303
    db.rollback()
    db.expire_all()  # see what the app committed
    db.refresh(cashier)
    assert cashier.can_receive_stock and not cashier.can_see_cost


def test_sidebar_badges(client, db, owner):
    p = seed_catalog(db, owner.id)
    p["sand"].min_stock_qty = 5000
    db.commit()
    login(client, owner, OWNER_PIN)
    assert 'class="nav-badge">1<' in client.get("/").text
