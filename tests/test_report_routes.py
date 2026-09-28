"""Phase 5 pages: reports, daily close and backup — permissions and rendering."""

import io

import pytest
from openpyxl import load_workbook

from app.models import DailyClose, Sale
from app.services import auth, backup
from app.services.sales import CartLineInput, CheckoutInput, create_sale
from tests.conftest import CASHIER_PIN, OWNER_PIN, login
from tests.factories import seed_catalog

REPORT_URLS = ["/reports/sales", "/reports/sales?group=month", "/reports/breakdown?by=product",
               "/reports/breakdown?by=category", "/reports/breakdown?by=cashier", "/reports/top?by=qty",
               "/reports/idle?days=30", "/reports/tax"]


@pytest.fixture
def sold(db, owner):
    p = seed_catalog(db, owner.id)
    p["tiger50"].avg_cost = 11600
    db.commit()
    bag = next(u.id for u in p["tiger50"].units if u.unit_name == "ถุง")
    sale = create_sale(db, owner, CheckoutInput(lines=[CartLineInput(p["tiger50"].id, bag, "3")]))
    db.commit()
    return sale


@pytest.mark.parametrize("url", REPORT_URLS)
def test_reports_render_for_owner(client, owner, sold, url):
    login(client, owner, OWNER_PIN)
    r = client.get(url)
    assert r.status_code == 200
    assert "ส่งออก Excel" in r.text


def test_breakdown_shows_profit(client, owner, sold):
    login(client, owner, OWNER_PIN)
    r = client.get("/reports/breakdown")
    assert "กำไรขั้นต้น" in r.text
    assert sold.lines[0].product_name_snapshot in r.text


@pytest.mark.parametrize("url", ["/reports/sales", "/reports/tax", "/reports/idle"])
def test_reports_export_xlsx(client, owner, sold, url):
    login(client, owner, OWNER_PIN)
    r = client.get(url + "?format=xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.active.max_row >= 3


def test_reports_forbidden_for_cashier(client, cashier):
    login(client, cashier, CASHIER_PIN)
    for url in ("/reports/sales", "/reports/breakdown", "/reports/sales?format=xlsx"):
        assert client.get(url).status_code == 403
    assert "/reports" not in client.get("/").text


def test_reports_allowed_for_cost_viewer(client, db, owner, cashier):
    auth.set_permissions(db, owner, cashier, can_receive_stock=False, can_see_cost=True)
    db.commit()
    login(client, cashier, CASHIER_PIN)
    assert client.get("/reports/breakdown").status_code == 200


def test_bad_dates_fall_back(client, owner):
    login(client, owner, OWNER_PIN)
    assert client.get("/reports/sales?start=xx&end=2026-13-40").status_code == 200


def test_cashier_can_close_day(client, db, cashier, sold):
    login(client, cashier, CASHIER_PIN)
    page = client.get("/close")
    assert page.status_code == 200 and "นับเงินในลิ้นชัก" in page.text
    r = client.post("/close", data={"opening": "1,000", "counted": "1500", "note": "ทดสอบ"}, follow_redirects=False)
    assert r.status_code == 303
    db.rollback()
    db.expire_all()
    c = db.query(DailyClose).one()
    assert c.opening_cash == 100000 and c.counted_cash == 150000
    assert c.expected_cash == 100000 + sold.total
    page = client.get(r.headers["location"])
    assert "บันทึกปิดยอดแล้ว" in page.text and "ล่าสุด" in page.text
    assert 'value="1000.00"' in page.text or "1000" in page.text  # opening remembered for re-close


def test_close_rejects_bad_amount(client, cashier):
    login(client, cashier, CASHIER_PIN)
    r = client.post("/close", data={"opening": "0", "counted": "abc"})
    assert r.status_code == 400
    assert "alert-error" in r.text


def test_backup_page_owner_only(client, cashier):
    login(client, cashier, CASHIER_PIN)
    assert client.get("/backup").status_code == 403
    assert client.post("/backup/now", follow_redirects=False).status_code == 403


def test_backup_now_and_page(client, app, owner):
    login(client, owner, OWNER_PIN)
    r = client.post("/backup/now", follow_redirects=False)
    assert r.status_code == 303
    page = client.get(r.headers["location"])
    assert page.status_code == 200
    assert "กู้คืนไฟล์นี้" in page.text
    assert "เริ่มใช้งานจริง" in page.text


def test_restore_wrong_pin(client, app, db, owner):
    login(client, owner, OWNER_PIN)
    client.post("/backup/now")
    target = backup.backup_dir(db, app.state.config)
    name = backup.list_backups(target)[0].name
    r = client.post("/backup/restore", data={"filename": name, "owner_pin": "0000"})
    assert r.status_code == 400
    assert "alert-error" in r.text


def test_go_live_route(client, db, owner, sold):
    login(client, owner, OWNER_PIN)
    bad = client.post("/backup/go-live", data={"owner_pin": OWNER_PIN, "confirm": "ok"})
    assert bad.status_code == 400
    r = client.post("/backup/go-live", data={"owner_pin": OWNER_PIN, "confirm": backup.GO_LIVE_CONFIRM},
                    follow_redirects=False)
    assert r.status_code == 303
    db.rollback()
    db.expire_all()
    assert db.query(Sale).count() == 0
    page = client.get("/backup")
    assert "เริ่มใช้งานจริงแล้วเมื่อ" in page.text
