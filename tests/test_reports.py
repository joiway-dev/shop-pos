"""Report figures against a small fixture of known sales (SPEC phase 5)."""

import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.services import closing, customers, receivables, reports
from app.services.clock import now_local
from app.services.customers import CustomerInput
from app.services.excel import MONEY, QTY, Column, build_xlsx
from app.services.reports import allocate
from app.services.sales import BuyerInput, CartLineInput, CheckoutInput, create_sale, void_sale
from app.services.settings import get_shop_settings
from tests.conftest import OWNER_PIN
from tests.factories import seed_catalog

TODAY = now_local().date()
YESTERDAY = TODAY - timedelta(days=1)


def unit(p, name):
    return next(u.id for u in p.units if u.unit_name == name)


def sell(db, actor, lines, days_ago=0, **kw):
    db.commit()
    sale = create_sale(db, actor, CheckoutInput(lines=lines, **kw))
    sale.created_at = sale.created_at - timedelta(days=days_ago)
    db.commit()
    return sale


@pytest.fixture
def fx(db, owner):
    p = seed_catalog(db, owner.id)
    p["tiger50"].avg_cost, p["sand"].avg_cost, p["rebar12"].avg_cost = 11600, 37800, 15000
    db.commit()
    cust = customers.save_customer(db, owner, None, CustomerInput("ลูกค้าเชื่อ", address="x", credit_limit="99999"))
    db.commit()
    a = sell(db, owner, [CartLineInput(p["tiger50"].id, unit(p["tiger50"], "ถุง"), "10"),
                         CartLineInput(p["sand"].id, unit(p["sand"], "คิว"), "1")], days_ago=1, bill_discount="100")
    b = sell(db, owner, [CartLineInput(p["rebar12"].id, unit(p["rebar12"], "เส้น"), "2")], days_ago=1,
             payment_type="transfer")
    c = sell(db, owner, [CartLineInput(p["sand"].id, unit(p["sand"], "คิว"), "2")], payment_type="credit",
             customer_id=cust.id)
    d = sell(db, owner, [CartLineInput(p["tiger50"].id, unit(p["tiger50"], "ถุง"), "1")])
    void_sale(db, owner, d, "ทดสอบ", OWNER_PIN)
    db.commit()
    return {"p": p, "cust": cust, "a": a, "b": b, "c": c, "d": d}


@pytest.mark.parametrize("total,weights,expected", [
    (180000, [145000, 45000], [137368, 42632]),
    (100, [1, 1, 1], [34, 33, 33]),
    (0, [5, 5], [0, 0]),
    (7, [0, 0], [7, 0]),
])
def test_allocate(total, weights, expected):
    parts = allocate(total, weights)
    assert parts == expected and sum(parts) == total


def test_sales_by_day(db, fx):
    rows = reports.sales_by_period(db, YESTERDAY, TODAY)
    assert [r.key for r in rows] == [YESTERDAY.isoformat(), TODAY.isoformat()]
    y, t = rows
    assert (y.bills, y.gross, y.discount, y.cash, y.transfer, y.credit, y.voided) == (2, 217800, 10000, 180000, 37800, 0, 0)
    assert (t.bills, t.gross, t.credit, t.voided) == (1, 90000, 90000, 1)
    total = reports.total_row(rows)
    assert (total.bills, total.gross) == (3, 307800)


def test_sales_by_month(db, fx):
    rows = reports.sales_by_period(db, TODAY - timedelta(days=40), TODAY, "month")
    assert sum(r.gross for r in rows) == 307800


def test_product_breakdown_and_profit(db, fx):
    rows = {r.label: r for r in reports.breakdown(db, YESTERDAY, TODAY, "product")}
    tiger, sand, rebar = rows["ปูนซีเมนต์ตราเสือ 50 กก."], rows["ทรายหยาบ"], rows["เหล็กเส้นกลม 12 มม."]
    assert (tiger.qty, tiger.gross, tiger.cost, tiger.profit) == (10_000, 137368, 116000, 21368)
    assert (sand.qty, sand.gross, sand.cost) == (3_000, 42632 + 90000, 3 * 37800)
    assert (rebar.gross, rebar.cost, rebar.bill_count) == (37800, 30000, 1)
    total = reports.group_total(list(rows.values()))
    assert total.gross == 307800 and total.bill_count == 3
    assert "ปูนซีเมนต์ตราเสือ 50 กก." in rows and len(rows) == 3  # voided sale not counted


def test_category_and_cashier_breakdown(db, fx):
    cats = {r.label: r.gross for r in reports.breakdown(db, YESTERDAY, TODAY, "category")}
    assert sum(cats.values()) == 307800
    cashiers = reports.breakdown(db, YESTERDAY, TODAY, "cashier")
    assert [(r.label, r.gross) for r in cashiers] == [("เจ้าของ", 307800)]


def test_top_sellers_and_non_moving(db, fx):
    top = reports.top_sellers(db, YESTERDAY, TODAY)
    assert [r.label for r in top] == ["ปูนซีเมนต์ตราเสือ 50 กก.", "ทรายหยาบ", "เหล็กเส้นกลม 12 มม."]
    assert reports.top_sellers(db, YESTERDAY, TODAY, by="qty")[0].label == "ปูนซีเมนต์ตราเสือ 50 กก."
    idle = {r.product.name for r in reports.non_moving(db, days=60, today=TODAY)}
    assert idle == {"ปูนซีเมนต์ตราเสือ 40 กก.", "เหล็กเส้นกลม 9 มม.", "ท่อ PVC 4 นิ้ว", "สายไฟ VAF 2x1.5", "ตะปู 2 นิ้ว"}


def test_vat_on_profit_uses_net_amount(db, owner):
    p = seed_catalog(db, owner.id)
    s = get_shop_settings(db)
    s.vat_mode, s.tax_id, s.branch_no, s.allow_abbreviated_invoice = "vat", "0105555000005", "00000", True
    p["sand"].avg_cost = 30000
    db.commit()
    sell(db, owner, [CartLineInput(p["sand"].id, unit(p["sand"], "คิว"), "1")])  # 450 incl VAT
    row = reports.breakdown(db, TODAY, TODAY)[0]
    assert (row.gross, row.net, row.cost, row.profit) == (45000, 42056, 30000, 12056)


def test_tax_report(db, owner):
    p = seed_catalog(db, owner.id)
    sell(db, owner, [CartLineInput(p["sand"].id, unit(p["sand"], "คิว"), "1")])  # before VAT registration
    s = get_shop_settings(db)
    s.vat_mode, s.tax_id, s.branch_no, s.allow_abbreviated_invoice = "vat", "0105555000005", "00000", True
    db.commit()
    abb = sell(db, owner, [CartLineInput(p["sand"].id, unit(p["sand"], "คิว"), "1")])
    tax = sell(db, owner, [CartLineInput(p["tiger50"].id, unit(p["tiger50"], "ถุง"), "10")],
               buyer=BuyerInput("บริษัท ก จำกัด", "1 ถ.ก", "0105555000005", "3"))
    gone = sell(db, owner, [CartLineInput(p["sand"].id, unit(p["sand"], "คิว"), "1")])
    void_sale(db, owner, gone, "คีย์ผิด", OWNER_PIN)
    db.commit()
    rows = reports.tax_report(db, TODAY, TODAY)
    assert [r.sale.id for r in rows] == [abb.id, gone.id, tax.id]  # ABB... then INV...
    by_id = {r.sale.id: r for r in rows}
    assert by_id[abb.id].buyer_name.startswith("ลูกค้าทั่วไป")
    assert (by_id[tax.id].tax_id, by_id[tax.id].branch, by_id[tax.id].vat) == ("0105555000005", "สาขา 00003", 9486)
    assert by_id[gone.id].voided and by_id[gone.id].vat == 0 and by_id[gone.id].total == 0


def test_daily_close(db, owner, fx):
    receivables.receive_payment(db, owner, fx["cust"], "500", "cash")
    db.commit()
    sell(db, owner, [CartLineInput(fx["p"]["sand"].id, unit(fx["p"]["sand"], "คิว"), "1")])  # 450 cash today
    p = closing.preview(db, TODAY, 100000)
    assert (p["cash_sales"], p["cash_receipts"], p["credit"], p["expected"]) == (45000, 50000, 90000, 195000)
    row = closing.close_day(db, owner, TODAY, "1000", "1,940", "จ่ายค่าน้ำแข็ง 10 บาท")
    db.commit()
    assert (row.expected_cash, row.counted_cash, row.difference) == (195000, 194000, -1000)
    closing.close_day(db, owner, TODAY, "1000", "1950")
    db.commit()
    latest = closing.closes_of(db, TODAY)
    assert len(latest) == 2 and latest[0].difference == 0
    with pytest.raises(closing.CloseError):
        closing.close_day(db, owner, TODAY, "x", "1")


def test_build_xlsx_numbers():
    content = build_xlsx("ยอดขาย", "ทดสอบ", [Column("วัน"), Column("จำนวน", QTY), Column("ยอด", MONEY)],
                         [["1", 2500, 145050]], total=["รวม", 2500, 145050])
    ws = load_workbook(io.BytesIO(content)).worksheets[0]
    assert ws["A1"].value == "ยอดขาย"
    assert [c.value for c in ws[5]] == ["1", 2.5, 1450.5]
    assert ws["C5"].number_format == "#,##0.00"
    assert ws["A6"].value == "รวม"
