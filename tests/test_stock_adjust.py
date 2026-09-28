import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import AuditLog, Product
from app.services import auth, documents, opening_import, product_import, stock
from app.services.settings import SettingsError, ShopSettingsInput, get_shop_settings, update_shop_settings
from tests.factories import make_product


@pytest.fixture
def sand(db, owner):
    p = make_product(db, owner.id, "ทรายหยาบ", "คิว", "450", sku="S1")
    stock.record_movement(db, p.id, 10_000, "adjust", owner.id)
    db.commit()
    return p


def test_count_sets_balance(db, owner, sand):
    m = stock.adjust_to_count(db, owner, sand, "7.5", "นับสต็อกสิ้นเดือน")
    db.commit()
    assert m.qty_base == -2500 and stock.balance_of(db, sand.id) == 7500
    assert db.scalar(select(func.count()).where(AuditLog.action == "stock_adjust")) == 1
    with pytest.raises(stock.StockError):
        stock.adjust_to_count(db, owner, sand, "7.5", "นับซ้ำ")


def test_write_off(db, owner, sand):
    stock.write_off(db, owner, sand, "2", "damaged", "ถุงแตก")
    db.commit()
    assert stock.balance_of(db, sand.id) == 8000
    assert "ชำรุด" in stock.movements_of(db, sand.id)[0][0].note


@pytest.mark.parametrize("call", [
    lambda db, u, p: stock.adjust_to_count(db, u, p, "5", "ไม่มีสิทธิ์"),
    lambda db, u, p: stock.write_off(db, u, p, "1", "lost", "ไม่มีสิทธิ์"),
])
def test_only_owner_adjusts(db, cashier, sand, call):
    with pytest.raises(stock.StockError):
        call(db, cashier, sand)


def test_adjust_validation(db, owner, sand):
    with pytest.raises(stock.StockError):
        stock.adjust_to_count(db, owner, sand, "5", "")
    with pytest.raises(stock.StockError):
        stock.adjust_to_count(db, owner, sand, "-1", "ติดลบ")
    with pytest.raises(stock.StockError):
        stock.write_off(db, owner, sand, "1", "stolen", "ประเภทผิด")


def test_stock_rows_low_stock_and_value(db, owner, sand):
    sand.min_stock_qty = 20_000
    sand.avg_cost = 30000
    other = make_product(db, owner.id, "ปูน", "ถุง", "145")
    db.commit()
    rows, total = stock.stock_rows(db)
    assert total == 2 and dict((p.id, b) for p, b in rows) == {sand.id: 10_000, other.id: 0}
    assert [p.id for p, _ in stock.stock_rows(db, low_only=True)[0]] == [sand.id]
    assert stock.low_stock_count(db) == 1
    assert stock.stock_value(db) == 300000


def test_movements_running_balance(db, owner, sand):
    stock.write_off(db, owner, sand, "3", "lost", "หาย")
    rows = stock.movements_of(db, sand.id)
    assert [bal for _, bal in rows] == [7000, 10_000]


# --- opening balances ----------------------------------------------------------


def opening_rows(*rows):
    return [{opening_import.COL_SKU: s, opening_import.COL_QTY: q, opening_import.COL_COST: c} for s, q, c in rows]


def test_opening_sets_balance_and_cost(db, owner, sand):
    report = opening_import.run_opening(db, owner, opening_rows(("S1", "25", "310")), commit=True)
    assert (report.changed, report.failed) == (1, 0)
    assert stock.balance_of(db, sand.id) == 25_000
    assert db.get(Product, sand.id).avg_cost == 31000


def test_opening_preview_rolls_back(db, owner, sand):
    report = opening_import.run_opening(db, owner, opening_rows(("S1", "25", "310")), commit=False)
    assert report.changed == 1 and not report.committed
    assert stock.balance_of(db, sand.id) == 10_000


def test_opening_errors_and_same(db, owner, sand):
    report = opening_import.run_opening(db, owner, opening_rows(("NOPE", "1", ""), ("S1", "10", ""), ("S1", "5", "")),
                                        commit=False)
    assert [r.action for r in report.rows] == ["error", "same", "error"]  # 3rd: duplicate SKU in file
    report = opening_import.run_opening(db, owner, opening_rows(("S1", "abc", "")), commit=False)
    assert report.rows[0].action == "error"


def test_opening_owner_only(db, cashier, sand):
    with pytest.raises(opening_import.OpeningError):
        opening_import.run_opening(db, cashier, opening_rows(("S1", "1", "1")), commit=False)


def test_opening_template_roundtrip(db, owner, sand):
    content = opening_import.template_xlsx(db)
    ws = load_workbook(io.BytesIO(content)).worksheets[0]
    assert [c.value for c in ws[2]][:4] == ["S1", "ทรายหยาบ", "คิว", "10"]
    rows = product_import.read_rows("x.xlsx", content, required=opening_import.REQUIRED)
    rows[0][opening_import.COL_QTY] = "12"
    assert opening_import.run_opening(db, owner, rows, commit=True).changed == 1
    assert stock.balance_of(db, sand.id) == 12_000


# --- document prefixes ---------------------------------------------------------


def test_clean_prefixes():
    prefixes, errors = documents.clean_prefixes({"RC": " sb ", "GR": ""})
    assert prefixes["RC"] == "SB" and prefixes["GR"] == "GR" and not errors
    _, errors = documents.clean_prefixes({"RC": "INV"})
    assert "RC" in errors or "TAX" in errors
    _, errors = documents.clean_prefixes({"RC": "ใบเสร็จ"})
    assert "RC" in errors


def test_prefix_setting_used_for_new_documents(db, owner):
    update_shop_settings(db, owner.id, ShopSettingsInput(shop_name="ร้าน", doc_prefixes={"RC": "SB"}))
    db.commit()
    s = get_shop_settings(db)
    assert documents.prefix_for(s, "RC") == "SB" and documents.prefix_for(s, "ABB") == "ABB"
    with pytest.raises(SettingsError) as e:
        update_shop_settings(db, owner.id, ShopSettingsInput(shop_name="ร้าน", doc_prefixes={"RC": "GR"}))
    assert any(k.startswith("prefix_") for k in e.value.errors)


def test_permissions_only_for_staff(db, owner, cashier):
    auth.set_permissions(db, owner, cashier, True, False)
    assert cashier.may_receive_stock and not cashier.may_see_cost
    with pytest.raises(auth.AuthError):
        auth.set_permissions(db, owner, owner, True, True)
    with pytest.raises(auth.AuthError):
        auth.set_permissions(db, cashier, cashier, True, True)
    assert db.scalar(select(func.count()).where(AuditLog.action == "user_permissions")) == 1
